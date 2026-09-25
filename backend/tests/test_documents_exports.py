import csv
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.config import Settings
from app.main import create_app


@pytest.fixture
def client_and_data(tmp_path):
    with TestClient(create_app(Settings(app_data_dir=tmp_path, _env_file=None))) as client:
        yield client, tmp_path


def create_application(client, company="Acme", title="Engineer", day="2026-09-18"):
    response = client.post("/api/applications", json={
        "job_title": title, "company": company, "date_applied": day,
        "job_url": f"https://example.com/{company.lower().replace(' ', '-')}",
    })
    assert response.status_code == 201
    return response.json()


def upload(client, application_id, filename="keiran-cv.pdf", content=b"test cv"):
    return client.post(
        f"/api/applications/{application_id}/documents",
        data={"document_type": "CV"},
        files={"file": (filename, content, "application/pdf")},
    )


def test_uploaded_document_belongs_to_application_and_is_stored(client_and_data):
    client, data_dir = client_and_data
    application = create_application(client)
    other = create_application(client, company="Other")

    response = upload(client, application["id"])
    assert response.status_code == 201
    document = response.json()
    assert document["application_id"] == application["id"]
    assert document["type"] == "CV"
    assert document["filename"] == "keiran-cv.pdf"
    assert document["external_reference"] is None
    stored = Path(document["storage_path"])
    assert stored.read_bytes() == b"test cv"
    assert stored.is_relative_to((data_dir / "documents" / application["id"]).resolve())
    assert client.get(f"/api/applications/{application['id']}/documents").json() == [document]
    assert client.get(f"/api/applications/{other['id']}/documents").json() == []
    download = client.get(f"/api/applications/{application['id']}/documents/{document['id']}/content")
    assert download.content == b"test cv"
    assert "keiran-cv.pdf" in download.headers["content-disposition"]


def test_reference_document_and_exclusive_source_validation(client_and_data):
    client, _ = client_and_data
    application = create_application(client)
    response = client.post(f"/api/applications/{application['id']}/documents", data={
        "document_type": "COVER_LETTER",
        "filename": "cover-letter.docx",
        "external_reference": "C:\\Documents\\cover-letter.docx",
    })
    assert response.status_code == 201
    assert response.json()["storage_path"] is None
    assert response.json()["external_reference"] == "C:\\Documents\\cover-letter.docx"
    assert client.get(f"/api/applications/{application['id']}/documents/{response.json()['id']}/content").status_code == 404
    assert client.post(f"/api/applications/{application['id']}/documents", data={"document_type": "CV"}).status_code == 422
    assert client.post(
        f"/api/applications/{application['id']}/documents",
        data={"document_type": "CV", "filename": "both.pdf", "external_reference": "ref"},
        files={"file": ("both.pdf", b"both", "application/pdf")},
    ).status_code == 422


def test_document_filename_filter(client_and_data):
    client, _ = client_and_data
    matching = create_application(client, company="Matching")
    create_application(client, company="Excluded")
    assert upload(client, matching["id"], "Keiran-Platform-CV.pdf").status_code == 201
    response = client.get("/api/applications", params={"document_filename": "platform-cv"})
    assert [item["id"] for item in response.json()] == [matching["id"]]


def test_csv_all_and_filtered_exports(client_and_data):
    client, _ = client_and_data
    alpha = create_application(client, company="Alpha", title="Backend Engineer", day="2026-09-20")
    create_application(client, company="Beta", title="Designer", day="2026-09-10")
    assert upload(client, alpha["id"], "alpha-cv.pdf").status_code == 201

    all_response = client.get("/api/exports/applications.csv")
    assert all_response.status_code == 200
    assert all_response.headers["content-disposition"] == 'attachment; filename="applications.csv"'
    all_rows = list(csv.DictReader(StringIO(all_response.content.decode("utf-8-sig"))))
    assert [row["Company"] for row in all_rows] == ["Alpha", "Beta"]
    assert all_rows[0]["Documents"] == "alpha-cv.pdf"

    filtered = client.get("/api/exports/applications.csv", params={"company": "beta"})
    rows = list(csv.DictReader(StringIO(filtered.content.decode("utf-8-sig"))))
    assert [row["Company"] for row in rows] == ["Beta"]


def test_xlsx_formatting_dates_colors_and_filtered_rows(client_and_data):
    client, _ = client_and_data
    create_application(client, company="Alpha", day="2026-09-20")
    create_application(client, company="Beta", day="2026-09-10")
    response = client.get("/api/exports/applications.xlsx", params={"company": "alpha"})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Applications"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:M2"
    assert [cell.value for cell in sheet[1]] == [
        "Date Applied", "Company", "Job Title", "Status", "Outcome", "Location",
        "Remote Policy", "Contract Type", "Source", "Job URL", "Email Reference",
        "Posting Status", "Documents",
    ]
    assert sheet.max_row == 2
    assert sheet["A2"].value.date() == date(2026, 9, 20)
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert sheet["B2"].value == "Alpha"
    assert sheet["D2"].fill.fill_type == "solid"
    assert sheet.column_dimensions["C"].width >= 20
    workbook.close()
