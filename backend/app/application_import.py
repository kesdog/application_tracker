"""Bounded spreadsheet parsing and transactional, reviewed imports."""
import csv
from datetime import datetime, date
import hashlib
from io import BytesIO, StringIO
import json
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import text, select

from app import applications
from app.applications import InvalidApplication
from app.models import Application, ImportBatch
from app.schemas import ApplicationCreate, ApplicationUpdate

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
ALIASES = {'date applied': 'date_applied', 'job title': 'job_title', 'job url': 'job_url',
           'email reference': 'email_reference', 'phone number': 'phone_number', 'contact type': 'contact_type',
           'contract type': 'contract_type', 'remote policy': 'remote_policy', 'contact name': 'contact_name',
           'contact email': 'contact_email', 'deadline type': 'deadline_kind', 'posting status': 'posting_status'}
FIELDS = sorted(set(ApplicationCreate.model_fields) - {'followup_customization', 'followup_instructions'} | {'status', 'outcome', 'posting_status'})


class DuplicateIndex:
    """One payload-free scan per import, rather than one full scan per row."""
    def __init__(self, session):
        self.records, self.titles, self.urls = {}, {}, {}
        for item in session.execute(select(Application.id, Application.company, Application.job_title,
                Application.date_applied, Application.job_url).where(Application.deleted_at.is_(None))):
            self.add(item)

    def add(self, item):
        self.records[item.id] = item
        key = (applications.normalized(item.company), applications.normalized(item.job_title))
        self.titles.setdefault(key, set()).add(item.id)
        url = applications.normalized_url(item.job_url)
        if url: self.urls.setdefault(url, set()).add(item.id)

    def find(self, data):
        key = (applications.normalized(data.company), applications.normalized(data.job_title))
        url = applications.normalized_url(data.job_url)
        title_ids, url_ids = self.titles.get(key, set()), self.urls.get(url, set())
        matches = []
        for record_id in title_ids | url_ids:
            item = self.records[record_id]
            reasons = (['same company and title'] if record_id in title_ids else []) + (['same job URL'] if record_id in url_ids else [])
            if abs((data.date_applied - item.date_applied).days) <= 30: reasons.append('application dates within 30 days')
            matches.append({'id': item.id, 'company': item.company, 'job_title': item.job_title, 'date_applied': item.date_applied, 'reasons': reasons})
        return sorted(matches, key=lambda item: (item['date_applied'], item['id']), reverse=True)


def parse(filename, content):
    if len(content) > MAX_BYTES:
        raise ValueError('Imports must be 10 MiB or smaller')
    if filename.lower().endswith('.csv'):
        rows = csv.reader(StringIO(content.decode('utf-8-sig')))
        result = []
        try:
            for row in rows:
                if len(result) > MAX_ROWS:
                    raise ValueError('Import at most 5,000 rows at a time')
                if len(row) > 100:
                    raise ValueError('Import at most 100 columns')
                result.append([value[1:] if value.startswith("'") and re.match(r'^[\s\ufeff]*[=+@-]', value[1:]) else value for value in row])
        except csv.Error:
            raise ValueError('The CSV contains an oversized or invalid field. Correct it before importing.') from None
    elif filename.lower().endswith('.xlsx'):
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('The expanded workbook exceeds 100 MiB')
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                result = []
                for row in workbook.active.iter_rows():
                    if len(result) > MAX_ROWS or len(row) > 100:
                        raise ValueError('Import at most 5,000 rows and 100 columns')
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Replace spreadsheet formulas with literal values before importing')
                    result.append([cell.value for cell in row])
            finally:
                workbook.close()
        except (BadZipFile, KeyError) as exc:
            raise ValueError('The workbook could not be read') from exc
    else:
        raise ValueError('Choose a CSV or XLSX file')
    if not result:
        raise ValueError('The spreadsheet is empty')
    headers = [str(value or '').strip() for value in result[0]]
    if len(set(headers)) != len(headers) or any(not value for value in headers):
        raise ValueError('Column headings must be nonempty and unique')
    return headers, result[1:]


def prepare(session, filename, content, mapping=None):
    headers, rows = parse(filename, content)
    if mapping is None:
        mapping = {header: ALIASES.get(header.casefold(), header.casefold().replace(' ', '_')) for header in headers}
        mapping = {header: field if field in FIELDS else '' for header, field in mapping.items()}
    mapped = [field for field in mapping.values() if field]
    if len(mapped) != len(set(mapped)) or set(mapped) - set(FIELDS):
        raise ValueError('Map each supported application field at most once')
    if set(mapping) - set(headers):
        raise ValueError('The mapping contains an unknown column')
    results = []
    duplicate_index = DuplicateIndex(session)
    for index, row in enumerate(rows, start=2):
        if not any(value not in ('', None) for value in row):
            continue
        raw = {mapping[header]: value for header, value in zip(headers, row)
               if mapping.get(header) and value not in ('', None)}
        for key, value in list(raw.items()):
            if isinstance(value, (date, datetime)):
                raw[key] = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
        lifecycle = {key: raw.pop(key) for key in ('status', 'outcome', 'posting_status') if key in raw}
        try:
            item = ApplicationCreate.model_validate(raw)
            updates = ApplicationUpdate.model_validate(lifecycle)
            if updates.outcome and updates.status and updates.status.value != 'CLOSED':
                raise ValueError('An outcome requires CLOSED status')
            matches = duplicate_index.find(item)
            results.append({'row': index, 'application': item.model_dump(mode='json'),
                'updates': updates.model_dump(mode='json', exclude_unset=True), 'errors': [], 'duplicates': matches})
        except (ValidationError, ValueError) as exc:
            errors = [f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()] if isinstance(exc, ValidationError) else [str(exc)]
            results.append({'row': index, 'application': raw, 'updates': {}, 'errors': errors, 'duplicates': []})
    return {'headers': headers, 'fields': FIELDS, 'mapping': {h: mapping.get(h) if mapping.get(h) in FIELDS else '' for h in headers}, 'rows': results}


def apply(session, settings, filename, content, mapping, selected, duplicate_mode):
    if duplicate_mode not in ('skip', 'create'):
        raise ValueError('Choose skip or create for duplicates')
    if not selected or len(selected) != len(set(selected)):
        raise ValueError('Select each row at most once, and select at least one row')
    digest = hashlib.sha256(content + json.dumps([mapping, sorted(selected), duplicate_mode], sort_keys=True).encode()).hexdigest()
    session.execute(text('BEGIN IMMEDIATE'))
    receipt = session.get(ImportBatch, digest)
    if receipt:
        return receipt.result
    preview = prepare(session, filename, content, mapping)
    rows = {row['row']: row for row in preview['rows']}
    if set(selected) - rows.keys() or any(rows[index]['errors'] for index in selected):
        raise ValueError('Review and correct the selected rows before importing')
    result = {'created': [], 'skipped': 0}
    duplicate_index = DuplicateIndex(session)
    for index in sorted(selected):
        row = rows[index]
        item = ApplicationCreate.model_validate(row['application'])
        matches = duplicate_index.find(item)
        if duplicate_mode == 'skip' and matches:
            result['skipped'] += 1
            continue
        parent = applications.create_application(session, item, settings=settings, commit=False, possible_matches=matches)
        if row['updates']:
            applications.update_application(session, parent.id, ApplicationUpdate(**row['updates']), settings=settings, commit=False)
        result['created'].append(parent.id)
        duplicate_index.add(parent)
    session.add(ImportBatch(id=digest, result=result))
    session.commit()
    return result
