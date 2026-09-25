import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import applications
from app.models import ApplicationDocument
from app.schemas import ApplicationFilters


HEADERS = [
    "Date Applied", "Company", "Job Title", "Status", "Outcome", "Location",
    "Remote Policy", "Contract Type", "Source", "Job URL", "Email Reference",
    "Phone Number", "Posting Status", "Documents",
]


def _rows(session: Session, filters: ApplicationFilters) -> list[list]:
    records = applications.list_applications(session, filters)
    document_rows = session.execute(select(ApplicationDocument.application_id, ApplicationDocument.filename)).all()
    documents: dict[str, list[str]] = {}
    for application_id, filename in document_rows:
        documents.setdefault(application_id, []).append(filename)
    return [[
        record.date_applied, record.company, record.job_title, record.status.value,
        record.outcome.value if record.outcome else "", record.location or "",
        record.remote_policy or "", record.contract_type or "", record.source or "",
        record.job_url or "", record.email_reference or "", record.phone_number or "", record.posting_status.value,
        "; ".join(sorted(documents.get(record.id, []), key=str.casefold)),
    ] for record in records]


def csv_export(session: Session, filters: ApplicationFilters) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    for row in _rows(session, filters):
        writer.writerow([value.isoformat() if hasattr(value, "isoformat") else value for value in row])
    return output.getvalue().encode("utf-8-sig")


def xlsx_export(session: Session, filters: ApplicationFilters) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(HEADERS)
    for row in _rows(session, filters):
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="263E5C")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{max(sheet.max_row, 1)}"
    sheet.column_dimensions["A"].width = 14
    for column in ("B", "C"):
        sheet.column_dimensions[column].width = 28
    for column in ("D", "E", "F", "G", "H", "I", "L", "M"):
        sheet.column_dimensions[column].width = 18
    for column in ("J", "K", "N"):
        sheet.column_dimensions[column].width = 36

    fills = {
        "SUBMITTED": "DCE9F9", "INTERVIEW": "E7DCF4", "CLOSED": "E2E6EA",
        "SUCCESSFUL": "DCEFE3", "UNSUCCESSFUL": "F7DDE0", "WITHDRAWN": "E2E6EA",
        "JOB_CANCELLED": "FBE7C9", "GHOSTED": "E2E6EA",
    }
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "yyyy-mm-dd"
        for column in (4, 5):
            value = sheet.cell(row, column).value
            if value in fills:
                sheet.cell(row, column).fill = PatternFill("solid", fgColor=fills[value])
        for cell in sheet[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
